"""
Wire reserved-word paths in DFA_States.xlsx so IF/THEN/ELSE/END/REPEAT/UNTIL/READ/WRITE
each end in an accept state. Adds columns l,s,d and states q40–q42 where needed.
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "DFA_States.xlsx"


def col_map(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def row_for(ws, name):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == name:
            return r
    return None


def main():
    wb = openpyxl.load_workbook(PATH)
    ws = wb.active

    cmap = col_map(ws)
    if "l" not in cmap:
        ws.insert_cols(16, 3)
        ws.cell(1, 16, "l")
        ws.cell(1, 17, "s")
        ws.cell(1, 18, "d")
        cmap = col_map(ws)
        for r in range(2, ws.max_row + 1):
            for c in (16, 17, 18):
                ws.cell(r, c, "qd")

    cmap = col_map(ws)

    def setv(state, colname, val):
        r, c = row_for(ws, state), cmap[colname]
        if r is None:
            raise SystemExit(f"missing row {state}")
        ws.cell(r, c, val)

    # Existing states: keyword edges
    setv("q21", "n", "q8")  # then → q8
    setv("q37", "t", "q12")  # repeat → q12
    setv("q36", "e", "q23")  # write → q23
    setv("q30", "l", "q14")  # until → q14
    setv("q25", "d", "q19")  # read → q19
    setv("q4", "l", "q40")  # else
    setv("q4", "n", "q42")  # end

    # New states for else / end prefixes
    for q in ("q40", "q41", "q42"):
        if row_for(ws, q) is None:
            row = [q]
            for c in range(2, ws.max_column + 1):
                h = ws.cell(1, c).value
                if h == "is_final":
                    row.append("No")
                else:
                    row.append("qd")
            ws.append(row)

    setv("q40", "s", "q41")
    setv("q41", "e", "q9")
    setv("q42", "d", "q11")

    wb.save(PATH)
    print("Updated", PATH)


if __name__ == "__main__":
    main()
